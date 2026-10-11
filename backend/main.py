from __future__ import annotations

import base64
import json
import os
import re
import shutil
import subprocess
import uuid
import zipfile
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import httpx
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
RENDS = DATA / "renders"
for folder in (DATA, UPLOADS, RENDS):
    folder.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="LessonLens 教案点评系统", version="0.1.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
EMU_PER_INCH = 914400
RENDER_W, RENDER_H = 1280, 720


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def safe_name(filename: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_.-]", "_", filename)


def extract_slides(path: Path) -> list[dict[str, Any]]:
    """Extract slide text and basic shape metrics without requiring Office."""
    slides: list[dict[str, Any]] = []
    try:
        prs = Presentation(str(path))
        for idx, slide in enumerate(prs.slides, 1):
            texts: list[str] = []
            shape_count = 0
            image_count = 0
            for shape in slide.shapes:
                shape_count += 1
                if getattr(shape, "shape_type", None) == 13:
                    image_count += 1
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            slides.append({"number": idx, "title": texts[0][:90] if texts else f"第 {idx} 页", "text": "\n".join(texts), "shape_count": shape_count, "image_count": image_count})
    except Exception:
        if path.suffix.lower() in {".pptx"}:
            with zipfile.ZipFile(path) as zf:
                names = sorted((n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)), key=lambda n: int(re.search(r"(\d+)", n).group(1)))
                for idx, name in enumerate(names, 1):
                    root = ET.fromstring(zf.read(name))
                    texts = [node.text or "" for node in root.findall(".//a:t", NS)]
                    slides.append({"number": idx, "title": texts[0][:90] if texts else f"第 {idx} 页", "text": "\n".join(texts), "shape_count": len(texts), "image_count": 0})
    return slides


def normalize_source(path: Path, job_id: str) -> Path:
    """Convert legacy .ppt files when LibreOffice is available."""
    if path.suffix.lower() != ".ppt":
        return path
    out = UPLOADS / job_id
    out.mkdir(exist_ok=True)
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pptx", "--outdir", str(out), str(path)], check=True, timeout=45, capture_output=True)
        converted = out / f"{path.stem}.pptx"
        if converted.exists():
            return converted
    except (FileNotFoundError, subprocess.SubprocessError):
        pass
    return path


def render_with_libreoffice(path: Path, job_id: str) -> list[str]:
    out = RENDS / job_id
    out.mkdir(exist_ok=True)
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)], check=True, timeout=45, capture_output=True)
        pdf = out / f"{path.stem}.pdf"
        if pdf.exists():
            subprocess.run(["pdftoppm", "-png", "-r", "110", str(pdf), str(out / "slide")], check=True, timeout=45, capture_output=True)
    except (FileNotFoundError, subprocess.SubprocessError):
        return []
    return [p.name for p in sorted(out.glob("slide-*.png"))]


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size)
            except OSError:
                continue
    return ImageFont.load_default()


def _rgb(value: Any, fallback: tuple[int, int, int]) -> tuple[int, int, int]:
    try:
        return tuple(value)[:3] if value else fallback
    except Exception:
        return fallback


def _shape_fill(shape: Any, fallback: tuple[int, int, int]) -> tuple[int, int, int]:
    try:
        if shape.fill.type is not None and shape.fill.fore_color.type is not None:
            return _rgb(shape.fill.fore_color.rgb, fallback)
    except Exception:
        pass
    return fallback


def _wrap_lines(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.replace("\r", "").split("\n"):
        current = ""
        for char in paragraph:
            candidate = current + char
            if current and draw.textlength(candidate, font=font) > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        lines.append(current)
    return lines or [""]


def render_with_python(path: Path, job_id: str) -> list[str]:
    """Render a readable slide image when an Office renderer is unavailable."""
    out = RENDS / job_id
    out.mkdir(exist_ok=True)
    try:
        prs = Presentation(str(path))
        sx, sy = RENDER_W / prs.slide_width, RENDER_H / prs.slide_height
        files: list[str] = []
        for index, slide in enumerate(prs.slides, 1):
            canvas = Image.new("RGB", (RENDER_W, RENDER_H), (255, 255, 255))
            draw = ImageDraw.Draw(canvas)
            try:
                canvas.paste(_rgb(slide.background.fill.fore_color.rgb, (255, 255, 255)), (0, 0, RENDER_W, RENDER_H))
            except Exception:
                pass
            for shape in slide.shapes:
                left, top = int(shape.left * sx), int(shape.top * sy)
                right, bottom = int((shape.left + shape.width) * sx), int((shape.top + shape.height) * sy)
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        source = Image.open(BytesIO(shape.image.blob)).convert("RGBA")
                        image = Image.new("RGBA", source.size, (255, 255, 255, 255))
                        image.alpha_composite(source)
                        image = image.convert("RGB")
                        image.thumbnail((max(1, right - left), max(1, bottom - top)), Image.Resampling.LANCZOS)
                        canvas.paste(image, (left, top))
                    except Exception:
                        pass
                    continue
                if getattr(shape, "shape_type", None) in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM}:
                    draw.rounded_rectangle((left, top, right, bottom), radius=4, fill=_shape_fill(shape, (242, 246, 246)), outline=(210, 220, 220))
                if hasattr(shape, "text") and shape.text.strip():
                    text = shape.text.strip()
                    font_size = max(12, min(42, int((shape.height / EMU_PER_INCH) * 72 * 0.78)))
                    font = _font(font_size, bold=font_size >= 24)
                    color = (35, 47, 53)
                    try:
                        if shape.text_frame.paragraphs and shape.text_frame.paragraphs[0].runs:
                            run_color = shape.text_frame.paragraphs[0].runs[0].font.color.rgb
                            color = _rgb(run_color, color)
                    except Exception:
                        pass
                    lines = _wrap_lines(draw, text, font, max(80, right - left - 18))
                    line_height = font_size + 6
                    for line_index, line in enumerate(lines[: max(1, int((bottom - top) / line_height))]):
                        draw.text((left + 9, top + 7 + line_index * line_height), line, font=font, fill=color)
            draw.rectangle((0, RENDER_H - 24, RENDER_W, RENDER_H), fill=(247, 249, 249))
            draw.text((24, RENDER_H - 19), f"{index:02d}  ·  LessonLens 结构化预览", font=_font(11), fill=(125, 140, 144))
            name = f"slide-{index:03d}.png"
            canvas.save(out / name, optimize=True)
            files.append(name)
        return files
    except Exception:
        return []


def _number(value: Any, default: int) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def normalize_review(review: Any, slides: list[dict[str, Any]]) -> dict[str, Any]:
    """Normalize model responses from different JSON naming conventions."""
    raw = review if isinstance(review, dict) else {}
    raw_dimensions = raw.get("dimensions") or raw.get("dimension_scores") or []
    if isinstance(raw_dimensions, dict):
        raw_dimensions = [{"name": key, "score": value} for key, value in raw_dimensions.items()]
    dimensions = []
    for item in raw_dimensions if isinstance(raw_dimensions, list) else []:
        if isinstance(item, dict):
            label = item.get("label") or item.get("name") or item.get("dimension") or item.get("title") or "观察维度"
            score = max(0, min(100, _number(item.get("score", item.get("value", item.get("rating", 0))), 0)))
            dimensions.append({"label": str(label), "score": score, "comment": item.get("comment", "")})
        elif isinstance(item, (int, float)):
            dimensions.append({"label": f"维度 {len(dimensions) + 1}", "score": max(0, min(100, int(item))), "comment": ""})
    if not dimensions:
        base = _number(raw.get("overall_score"), 0)
        dimensions = [{"label": name, "score": max(0, min(100, base)), "comment": ""} for name in ("内容逻辑", "视觉设计", "课堂可用性")]

    raw_pages = raw.get("page_reviews") or raw.get("pages") or raw.get("slide_reviews") or []
    pages = []
    for index, item in enumerate(raw_pages if isinstance(raw_pages, list) else []):
        item = item if isinstance(item, dict) else {}
        number = _number(item.get("page", item.get("number", item.get("page_number", index + 1))), index + 1)
        title = item.get("headline") or item.get("title") or item.get("name") or f"第 {number} 页"
        content_logic = item.get("content_logic") or item.get("content") or ""
        visual_design = item.get("visual_design") or item.get("visual") or ""
        classroom = item.get("classroom_usability") or item.get("usability") or ""
        suggestion = item.get("suggestions") or item.get("suggestion") or item.get("recommendation") or []
        if isinstance(suggestion, str): suggestion = [suggestion]
        suggestion = [str(value) for value in suggestion if value]
        slide = next((value for value in slides if _number(value.get("number"), 0) == number), {})
        current_summary = item.get("summary")
        if current_summary in {"暂无文字点评", "暂无点评", "无文字点评"} or (isinstance(current_summary, str) and current_summary.startswith("本页重点：")):
            current_summary = ""
        summary = current_summary or content_logic or visual_design or classroom or item.get("feedback") or item.get("observation") or item.get("comment") or item.get("critique") or item.get("commentary") or item.get("explanation") or item.get("analysis") or item.get("review")
        if not summary and suggestion:
            slide_title = str(slide.get("title") or title).strip()
            summary = f"页面围绕“{slide_title[:46]}”展开。点评建议：{suggestion[0]}"
        if not summary:
            slide_text = str(slide.get("text") or slide.get("title") or "").strip()
            summary = f"页面包含：{slide_text[:110]}" if slide_text else "本页未提取到可读文字，建议结合预览检查版式与教学动作。"
        detail_parts = [part for part in (content_logic, visual_design, classroom) if part and part != summary]
        score = _number(item.get("score", item.get("overall_score", item.get("rating", 0))), 0)
        average_dimension_score = round(sum(d["score"] for d in dimensions) / len(dimensions))
        legacy_fallback_score = score == average_dimension_score and not any((content_logic, visual_design, classroom, item.get("details")))
        if not score or legacy_fallback_score:
            text_length = len(str(slide.get("text") or ""))
            penalty = (10 if text_length == 0 else 0) + (6 if text_length > 500 else 0) + (4 if slide.get("image_count", 0) == 0 and slide.get("shape_count", 0) <= 1 else 0)
            score = max(0, min(100, average_dimension_score - penalty))
        pages.append({"page": number, "headline": str(title), "summary": str(summary), "details": detail_parts, "tags": item.get("tags") if isinstance(item.get("tags"), list) else [], "score": score, "suggestions": suggestion})

    overall = _number(raw.get("overall_score"), 0) or (round(sum(item["score"] for item in pages) / len(pages), 1) if pages else round(sum(d["score"] for d in dimensions) / len(dimensions), 1))
    raw_actions = raw.get("actions") or raw.get("action_items") or []
    actions = []
    priority_map = {"高": "high", "中": "medium", "低": "low", "high": "high", "medium": "medium", "low": "low"}
    for item in raw_actions if isinstance(raw_actions, list) else []:
        item = item if isinstance(item, dict) else {"title": str(item)}
        priority = priority_map.get(str(item.get("priority", "medium")).lower(), "medium")
        pages_value = item.get("pages") or item.get("target_pages") or item.get("page_numbers") or []
        if not isinstance(pages_value, list): pages_value = [pages_value]
        actions.append({"priority": priority, "title": str(item.get("title") or item.get("action") or item.get("name") or "待处理建议"), "detail": str(item.get("detail") or item.get("description") or item.get("expected_effect") or ""), "pages": [_number(value, 0) for value in pages_value if _number(value, 0) > 0]})
    return {"overall_score": overall, "grade": raw.get("grade") or ("A-" if overall >= 85 else "B+" if overall >= 75 else "C+"), "confidence": raw.get("confidence", 0.88), "overview": raw.get("overview") or raw.get("summary") or "暂无总结", "dimensions": dimensions, "page_reviews": pages, "actions": actions}


def normalize_job(item: dict[str, Any]) -> dict[str, Any]:
    item["review"] = normalize_review(item.get("review"), item.get("slides", []))
    needs_preview = not item.get("preview_files") or (item.get("render_mode") == "python" and item.get("preview_version") != 2)
    if needs_preview:
        candidates = sorted(UPLOADS.glob(f"{item.get('id', '')}_*"))
        if candidates:
            source = candidates[0]
            previews = render_with_libreoffice(source, item["id"])
            mode = "office" if previews else "text"
            if not previews:
                previews = render_with_python(source, item["id"])
                mode = "python" if previews else "text"
            item["preview_files"] = previews
            item["rendered"] = bool(previews)
            item["render_mode"] = mode
            item["preview_version"] = 2
    return item


def demo_review(slides: list[dict[str, Any]]) -> dict[str, Any]:
    page_reviews = []
    for slide in slides:
        txt = slide["text"]
        has_title = bool(txt.strip())
        density = min(100, max(35, len(txt) // 2 + slide["shape_count"] * 2))
        score = max(68, 92 - (density - 55) // 4)
        page_reviews.append({
            "page": slide["number"], "score": score,
            "headline": "层级清晰，可继续强化课堂动作" if has_title else "页面缺少可识别的教学锚点",
            "summary": "核心信息已经出现，建议把‘看懂’进一步转成‘做出来’。" if has_title else "建议补充一句结论型标题，并明确本页要学生完成的动作。",
            "tags": (["结构清楚", "信息密度"] if density < 70 else ["信息密度偏高", "认知负担"]),
            "signals": {"clarity": max(60, min(96, score + 2)), "visual": max(58, score - 1), "pedagogy": max(63, score + 4)},
            "suggestions": ["标题改成结论句，让学生先知道为什么重要。", "在页脚加入一个 30 秒检查点或同伴讨论提示。"],
        })
    avg = round(sum(item["score"] for item in page_reviews) / len(page_reviews), 1) if page_reviews else 0
    return {"overall_score": avg, "grade": "A-" if avg >= 85 else "B+", "confidence": 0.88, "overview": "这份教案的叙事骨架已经成立，下一步是把内容节奏和课堂动作绑定得更紧。", "dimensions": [{"label": "内容逻辑", "score": min(96, round(avg + 2))}, {"label": "视觉表达", "score": max(0, round(avg - 3))}, {"label": "教学动作", "score": min(95, round(avg + 4))}], "page_reviews": page_reviews, "actions": [{"priority": "high", "title": "压缩第 3 页信息密度", "detail": "将三段说明改成 3 个要点，并增加一个判断题。", "pages": [3]}, {"priority": "medium", "title": "补齐课堂检查点", "detail": "每 4-5 页插入一次可观察的学生产出。", "pages": [2, 5]}, {"priority": "low", "title": "统一图表标注", "detail": "统一字号和图例位置，减少页面之间的视觉跳动。", "pages": [1, 4]}]}


def read_jobs() -> list[dict[str, Any]]:
    p = DATA / "jobs.json"
    if not p.exists(): return []
    return json.loads(p.read_text(encoding="utf-8"))


def write_jobs(items: list[dict[str, Any]]) -> None:
    (DATA / "jobs.json").write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


async def call_llm(slides: list[dict[str, Any]], preview_files: list[str], job_id: str) -> dict[str, Any] | None:
    api_key = os.getenv("LLM_API_KEY")
    base_url = os.getenv("LLM_BASE_URL", "https://api.openai.com/v1")
    model = os.getenv("LLM_MODEL", "gpt-4o-mini")
    if not api_key: return None
    prompt = "你是资深教师与教学设计师。请只返回 JSON，字段必须包含 overall_score, grade, overview, dimensions, page_reviews, actions。dimensions 的每项必须是 {label, score, comment}；page_reviews 的每项必须是 {page, headline, summary, tags, score, suggestions}，summary 必须写成完整的页面观察，不能留空或写‘暂无文字点评’。actions 的每项必须是 {priority, title, detail, pages}。逐页评价以下幻灯片的内容逻辑、视觉设计、课堂可用性，并给出可执行建议。页面图片按顺序附在对应页面的文字后面。\n" + json.dumps(slides, ensure_ascii=False)
    content: list[dict[str, Any]] = [{"type": "text", "text": prompt}]
    for index, slide in enumerate(slides):
        content.append({"type": "text", "text": f"第 {slide['number']} 页：{slide.get('text', '')}"})
        if index < len(preview_files):
            image_path = RENDS / job_id / preview_files[index]
            if image_path.exists():
                encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
                content.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{encoded}"}})
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    body = {"model": model, "temperature": 0.2, "response_format": {"type": "json_object"}, "messages": [{"role": "user", "content": content}]}
    try:
        async with httpx.AsyncClient(timeout=90) as client:
            response = await client.post(f"{base_url.rstrip('/')}/chat/completions", headers=headers, json=body)
            response.raise_for_status()
            content = response.json()["choices"][0]["message"]["content"]
            return json.loads(content)
    except Exception:
        return None


@app.get("/api/health")
async def health(): return {"ok": True, "llm_configured": bool(os.getenv("LLM_API_KEY"))}


@app.get("/api/jobs")
async def jobs():
    items = read_jobs()
    normalized = [normalize_job(item) for item in items]
    if normalized:
        write_jobs(normalized)
    return sorted(normalized, key=lambda x: x["created_at"], reverse=True)


@app.get("/api/jobs/{job_id}")
async def job(job_id: str):
    items = read_jobs()
    found = next((item for item in items if item["id"] == job_id), None)
    if not found: raise HTTPException(404, "Job not found")
    normalize_job(found)
    write_jobs(items)
    return found


@app.get("/api/jobs/{job_id}/slides/{slide_name}")
async def slide_image(job_id: str, slide_name: str):
    path = RENDS / job_id / safe_name(slide_name)
    if not path.exists(): raise HTTPException(404, "Preview not found")
    return FileResponse(path)


@app.post("/api/reviews")
async def create_review(file: UploadFile = File(...)):
    if not file.filename or Path(file.filename).suffix.lower() not in {".pptx", ".ppt"}:
        raise HTTPException(400, "请上传 .pptx 或 .ppt 文件")
    job_id = uuid.uuid4().hex[:10]
    source = UPLOADS / f"{job_id}_{safe_name(file.filename)}"
    with source.open("wb") as target: shutil.copyfileobj(file.file, target)
    parse_source = normalize_source(source, job_id)
    slides = extract_slides(parse_source)
    preview_files = render_with_libreoffice(parse_source, job_id)
    render_mode = "office" if preview_files else "text"
    if not preview_files:
        preview_files = render_with_python(parse_source, job_id)
        render_mode = "python" if preview_files else "text"
    if not slides:
        slides = [{"number": 1, "title": "无法解析的演示文稿", "text": "请确认文件未损坏，并重新导出为 PPTX。", "shape_count": 0, "image_count": 0}]
    llm = await call_llm(slides, preview_files, job_id)
    review = normalize_review(llm or demo_review(slides), slides)
    now = utc_now()
    item = {"id": job_id, "filename": file.filename, "created_at": now, "status": "completed", "rendered": bool(preview_files), "render_mode": render_mode, "preview_version": 2, "slides": slides, "preview_files": preview_files, "review": review}
    all_jobs = read_jobs(); all_jobs.append(item); write_jobs(all_jobs)
    return item


app.mount("/", StaticFiles(directory=ROOT / "frontend", html=True), name="frontend")
